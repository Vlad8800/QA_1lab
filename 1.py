import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.79)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.79)

normal_style = doc.styles['Normal']
normal_font = normal_style.font
normal_font.name = 'Times New Roman'
normal_font.size = Pt(14)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'<w:tcBorders %s/>' % nsdecls('w'))
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = parse_xml(r'<%s %s %s/>' % (tag, nsdecls('w'),
                                                   ' '.join([f'{k}="{v}"' for k, v in edge_data.items()])))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

# Титулка
p_inst = doc.add_paragraph()
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_inst.paragraph_format.line_spacing = 1.15
r = p_inst.add_run("Міністерство освіти і науки України\nЧернівецький національний університет імені Юрія Федьковича\nНавчально-науковий інститут фізико-технічних та комп'ютерних наук\nКафедра програмного забезпечення комп'ютерних систем")
r.bold = True
r.font.size = Pt(13)

doc.add_paragraph().paragraph_format.space_after = Pt(40)

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_rep = p_title.add_run("ЗВІТ\n")
r_rep.bold = True
r_rep.font.size = Pt(18)
r_sub = p_title.add_run("про виконання лабораторної роботи № 2\nз курсу «Безпека програм та даних»\n")
r_sub.bold = True
r_sub.font.size = Pt(14)
r_theme = p_title.add_run("Тема: Захист реєстру операційної системи Windows 11")
r_theme.bold = True
r_theme.font.size = Pt(14)

doc.add_paragraph().paragraph_format.space_after = Pt(90)

p_meta = doc.add_paragraph()
p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_meta.paragraph_format.line_spacing = 1.15
r_meta = p_meta.add_run("Виконав:\nстудент 4 курсу, групи 441\nЛазарєв В. В.\n\nПеревірив:\nд.ф.-м.н., проф. Остапов С. Е.")
r_meta.font.size = Pt(13)

doc.add_paragraph().paragraph_format.space_after = Pt(100)

p_city = doc.add_paragraph()
p_city.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_city = p_city.add_run("Чернівці — 2026")
r_city.bold = True
r_city.font.size = Pt(12)

doc.add_page_break()

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0, 51, 102)

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13.5)

def add_body(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)

def add_bullet(bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.add_run(bold_prefix).bold = True
    p.add_run(text)

add_heading_1("Мета роботи")
add_body("Ознайомитися з архітектурою та механізмами захисту реєстру операційної системи Windows 11. На практиці опанувати налаштування системного аудиту для критичних вуликів SAM та SECURITY, перевірку прав доступу до системного каталогу конфігурації Windows, конфігурування захисту віртуальної пам'яті при вимкненні ОС та блокування запису аварійних дампів для унеможливлення витоку автентифікаційних даних.")

add_heading_1("Хід роботи")
add_heading_2("1. Налаштування системної політики аудиту в secpol.msc")
add_body("В оснастці керування локальними політиками безпеки (secpol.msc) у розділі «Локальні політики» -> «Політика аудиту» активовано відстеження критичних подій:")
add_bullet("Audit logon events (Аудит входу в систему): ", "увімкнено фіксацію успішних (Success) та невдалих (Failure) спроб автентифікації.")
add_bullet("Audit object access (Аудит доступу до об'єктів): ", "увімкнено відстеження спроб звернення до захищених системних дескрипторів та ключів реєстру (Success, Failure).")

add_body("[Тут вставте скріншот вікна secpol.msc]")

add_heading_2("2. Налаштування списків аудиту (SACL) для вуликів SAM та SECURITY")
add_body("За допомогою редактора реєстру (regedit.exe) було налаштовано аудит звернень до гілок HKEY_LOCAL_MACHINE\\SAM та HKEY_LOCAL_MACHINE\\SECURITY:")
add_bullet("Вулик SAM: ", "у розширених параметрах безпеки на вкладці «Аудит» додано правило для групи Administrators (Тип: Усі, Застосовно до: Цей розділ і підрозділи, Дозволи: Повний доступ).")
add_bullet("Вулик SECURITY: ", "аналогічно сконфігуровано фіксацію всіх звернень для Administrators та SYSTEM.")

add_heading_2("3. Аудит та верифікація доступу до каталогу System32\\config")
add_body("У Провіднику відкрито властивості папки C:\\Windows\\System32\\config (вкладка «Безпека» -> «Додатково»). Перевірено відсутність небажаних сутностей:")
add_bullet("Результат перевірки: ", "звичайні користувачі (Users) та група Everyone позбавлені будь-якого доступу до файлів конфігурації; повний контроль надано виключно системним сутностям SYSTEM, TrustedInstaller, Administrators та власнику.")

add_body("[Тут вставте скріншот дозволів папки System32\\config]")

add_heading_2("4. Захист віртуальної пам'яті та вимкнення створення аварійних дампів")
add_body("З метою унеможливлення збереження залишкових хешів паролів та облікових даних у відкритому вигляді на накопичувачі реалізовано наступні захисні заходи:")
add_bullet("ClearPageFileAtShutdown: ", "у гілці HKLM\\SYSTEM\\CurrentControlSet\\Control\\Session Manager\\Memory Management встановлено прапорець ClearPageFileAtShutdown = 1 для обов'язкового перезапису сторінкового файлу нулями під час зупинки ОС.")
add_bullet("Вимкнення дампів пам'яті: ", "через «Додаткові параметри системи» в розділі «Запуск і відновлення» тип налагоджувальної інформації змінено на «(немає)», що блокує формування файлів memory.dmp.")

add_heading_1("Протокол зміни параметрів безпеки")
table = doc.add_table(rows=6, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Параметр / Об'єкт", "Значення до налаштування", "Встановлене значення", "Цільове призначення"]
widths = [Inches(1.8), Inches(1.5), Inches(1.5), Inches(2.2)]

for i, name in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = name
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].runs[0].font.size = Pt(11)
    set_cell_shading(cell, "EAECEE")

data = [
    ("Audit logon events", "Не визначено / Відімкнено", "Success, Failure", "Фіксація всіх спроб входу в систему"),
    ("Audit object access", "Не визначено / Відімкнено", "Success, Failure", "Активація підсистеми аудиту реєстру та файлів"),
    ("SACL вулика SAM", "Записи аудиту відсутні", "All / Full Control (Administrators)", "Виявлення спроб читання/викрадення бази SAM"),
    ("ClearPageFileAtShutdown", "0 (вимкнено)", "1 (увімкнено)", "Перезапис pagefile.sys нулями при вимкненні ОС"),
    ("Запис налагодження пам'яті", "Малий дамп пам'яті (256 КБ)", "(немає)", "Запобігання витоку паролів у файли *.dmp")
]

for row_idx, row_data in enumerate(data, start=1):
    for col_idx, text in enumerate(row_data):
        cell = table.rows[row_idx].cells[col_idx]
        cell.text = text
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if p.runs:
            p.runs[0].font.size = Pt(10.5)

border_kwargs = dict(
    top=dict(sz=4, val='single', color='D0D3D4'),
    bottom=dict(sz=4, val='single', color='D0D3D4'),
    left=dict(sz=4, val='single', color='D0D3D4'),
    right=dict(sz=4, val='single', color='D0D3D4')
)
for row in table.rows:
    for i, cell in enumerate(row.cells):
        cell.width = widths[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_border(cell, **border_kwargs)

add_heading_1("Висновок")
add_body("У ході виконання лабораторної роботи № 2 було детально досліджено архітектуру захисту реєстру та системних областей пам'яті ОС Windows 11. На практиці налаштовано механізми аудиту доступу до об'єктів (SACL) для критичних вуликів SAM та SECURITY, перевірено списки дозволів на каталог конфігурації System32\\config, а також реалізовано комплекс превентивних заходів проти атак викрадення залишкових облікових даних із файлу підкачування та аварійних дампів пам'яті. Результати конфігурування забезпечують захист системних паролів від несанкціонованого отримання.")

add_heading_1("Відповіді на контрольні запитання")
qa = [
    ("1. Які функції виконує реєстр Windows?",
     "Реєстр Windows є централізованою ієрархічною базою даних ОС, яка зберігає всі конфігураційні відомості про апаратне забезпечення, системні компоненти, служби, драйвери, параметри встановленого ПЗ та персональні налаштування облікових записів користувачів."),
    ("2. Які функції виконують вулики SAM та Security реєстру?",
     "Вулик SAM (Security Accounts Manager) містить облікові записи локальних користувачів та груп разом із незворотними NTLM-хешами їхніх паролів. Вулик Security зберігає глобальні параметри безпеки робочої станції: локальні політики, призначення прав і привілеїв, конфігурацію аудиту та системні секрети підсистеми LSA."),
    ("3. Які функції аудиту в системі безпеки операційних систем?",
     "Аудит забезпечує підзвітність дій у системі: він фіксує події входу в систему, спроби несанкціонованого звернення до захищених файлів або гілок реєстру, зміну прав доступу й аварійні події. Це дає змогу виявляти інциденти безпеки, проводити розслідування та відстежувати підозрілу активність."),
    ("4. Як можна отримати доступ до реєстру?",
     "Інтерактивно через графічні інтерфейси regedit.exe та regedt32.exe; за допомогою консольних утиліт reg.exe та командлетів PowerShell (Get-ItemProperty, Set-ItemProperty); а також програмно через виклики Win32 API (сімейство функцій RegOpenKeyEx, RegQueryValueEx тощо)."),
    ("5. Які основні принципи захисту реєстру Windows?",
     "Жорстке розмежування прав за допомогою списків DACL (блокування запису для звичайних користувачів); налаштування системного аудиту SACL на критичні гілки (SAM, SECURITY, Winlogon); обмеження або повне вимкнення служби Remote Registry; забезпечення фізичного захисту файлів вуликів у каталозі System32\\Config та контроль процесів з правами SYSTEM."),
    ("6. Які основні принципи захисту Windows?",
     "Ешелонований багаторівневий захист, реалізація принципу найменших привілеїв (робота під стандартним обліковим записом), обмеження фізичного доступу та парольний захист BIOS/UEFI з вимкненням зовнішнього завантаження, строгі політики паролів та блокування облікових записів, ізоляція адресного простору ядра від користувача та шифрування диска (BitLocker).")
]

for q, a in qa:
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(8)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.keep_with_next = True
    p_q.add_run(q).bold = True
    add_body(a)

doc.save("Звіт_Лабораторна_2_Реєстр_Лазарєв.docx")
print("Файл успішно створено!")